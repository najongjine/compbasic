from pathlib import Path
import urllib.request, json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).parent
A=ROOT/'assets'
def get(url,name):
    p=A/name
    if not p.exists():
        req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0'})
        with urllib.request.urlopen(req,timeout=40) as r: p.write_bytes(r.read())
    return p
get('https://cdn.loc.gov/service/pnp/cph/3b00000/3b07000/3b07400/3b07443r.jpg','columbus.jpg')
get('https://www.nps.gov/jame/learn/historyculture/images/D-4-Colonists-Landing-at-Jamestown-for-web.jpg','jamestown.jpg')
get('https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/ne_110m_land.geojson','land.json')
font='C:/Windows/Fonts/malgun.ttf'
bold='C:/Windows/Fonts/malgunbd.ttf'
def ft(s,b=False): return ImageFont.truetype(bold if b else font,s)
im=Image.new('RGB',(1500,830),'#edf5f8'); dr=ImageDraw.Draw(im)
def xy(lon,lat): return ((lon+105)/130*1500,(72-lat)/72*830)
geo=json.loads((A/'land.json').read_text())
for f in geo['features']:
    g=f['geometry']; polys=g['coordinates'] if g['type']=='MultiPolygon' else [g['coordinates']]
    for poly in polys: dr.polygon([xy(*v[:2]) for v in poly[0]],fill='#d5dccb',outline='#aab6a3')
def label(lon,lat,t): dr.text(xy(lon,lat),t,font=ft(30,True),fill='#243b43')
label(-99,52,'북아메리카'); label(-17,64,'유럽'); label(-15,22,'아프리카'); label(-53,29,'대서양')
def route(points,color):
    pts=[xy(*v) for v in points]; dr.line(pts,fill=color,width=6)
    x,y=pts[-1]; xx,yy=pts[-2]; a=math.atan2(y-yy,x-xx)
    dr.polygon([(x,y),(x-25*math.cos(a-.45),y-25*math.sin(a-.45)),(x-25*math.cos(a+.45),y-25*math.sin(a+.45))],fill=color)
route([(-6,37),(-17,28),(-38,24),(-59,24),(-74.5,24)],'#bb602d')
route([(-1,51),(-20,48),(-44,44),(-76.7,37.2)],'#2f668b')
dr.text((700,465),'1492  콜럼버스',font=ft(29,True),fill='#a44a21')
dr.text((585,225),'1607  제임스타운 정착',font=ft(27,True),fill='#245574')
for lon,lat,t,dx,dy in [(-6,37,'스페인',20,-15),(-1,51,'잉글랜드',15,-40),(-74.5,24,'바하마 일대',-135,15),(-76.7,37.2,'제임스타운',-175,-50)]:
    x,y=xy(lon,lat); dr.ellipse((x-7,y-7,x+7,y+7),fill='#283b40');dr.text((x+dx,y+dy),t,font=ft(25),fill='#243b43')
dr.text((35,774),'위치와 이동 방향을 보여 주는 개념도이며 실제 항로와 국경은 생략했습니다.',font=ft(23),fill='#52626c')
im.save(A/'atlantic.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=doc.styles[name];s.font.name='맑은 고딕';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕');s.font.color.rgb=RGBColor(0,0,0)
for s in doc.styles:
    for border in s.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    if s.type==1:
        sp=OxmlElement('w:snapToGrid');sp.set(qn('w:val'),'0');s.element.get_or_add_pPr().append(sp)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=Pt(16);normal.paragraph_format.space_after=Pt(8)
for name,size in [('Title',27),('Heading 1',21),('Heading 2',13)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_after=Pt(12)
doc.styles['Caption'].font.size=Pt(8.5)
doc.styles['Caption'].paragraph_format.space_after=Pt(12)
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(t):
    pp=doc.add_heading(t,1);pp.paragraph_format.page_break_before=True
def pic(name,cap,width=16.7):
    pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True;pp.paragraph_format.line_spacing=1.0
    r=pp.add_run()
    img=Image.open(A/name); width=min(width,7.0*img.width/img.height)
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r.add_picture(str(A/name),width=Cm(width))
    r._r.xpath('.//wp:docPr')[0].set('descr',cap)
    p(cap,'Caption')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFEAF0' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_before=Pt(6);pp.paragraph_format.space_after=Pt(6);pp.paragraph_format.line_spacing=1.15
                for r in pp.runs:r.font.size=Pt(9.5);r.bold=i==0
    return t

p('아메리카와 미국의 시작','Title')
p('유럽인의 도착에서 영국 식민지와 독립까지','Subtitle')
p('미국은 누군가가 발견한 뒤 곧바로 세운 나라가 아닙니다. 아메리카에는 이미 원주민 사회가 있었고, 유럽 국가들이 여러 지역을 식민지로 만들었습니다. 그 가운데 북아메리카 동부의 영국 식민지 13곳이 독립하면서 미국의 출발점이 되었습니다.')
pic('atlantic.png','그림 1  콜럼버스의 도착 지역과 영국인의 정착 지역은 달랐습니다. 지도 바탕: Natural Earth. 항로와 한글 표기는 설명용으로 제작.')
h('1  발견 이전에도 사람이 살고 있었습니다')
p('유럽인이 오기 오래전부터 아메리카에는 다양한 언어와 문화를 가진 원주민들이 살았습니다. 지역에 따라 농사를 짓고, 사냥과 어로를 하며, 교역망과 정치 공동체를 이루었습니다. 따라서 “신대륙 발견”은 주로 유럽인의 시각에서 붙인 표현입니다. [1]')
p('약 1000년 무렵에는 북유럽의 노르드인도 오늘날 캐나다 뉴펀들랜드에 도착했습니다. 랑스 오 메도 유적이 그 증거입니다. 다만 이 정착이 훗날 영국 식민지나 미국으로 이어진 것은 아닙니다. [2]')
p('용어 구분  아메리카는 남북아메리카를 포함하는 지리적 이름이고, 미국은 그 일부에 세워진 국가입니다. 1492년에 미국이라는 나라는 아직 없었습니다.')

page('2  콜럼버스의 항해가 바꾼 것')
p('1492년 콜럼버스는 스페인 왕실의 지원을 받아 대서양을 건넜습니다. 목적은 서쪽으로 항해해 아시아로 가는 길을 찾는 것이었습니다. 그가 도착한 곳은 오늘날 바하마 일대였고, 그곳에서 타이노 원주민을 만났습니다. 오늘날 미국 본토에 상륙한 사건으로 이해하면 안 됩니다. [3]')
pic('columbus.jpg','그림 2  콜럼버스의 1492년 히스파니올라 도착을 묘사한 후대의 삽화. 실제 현장을 그대로 기록한 그림이 아니며, 유럽 중심의 시각을 살펴볼 필요가 있습니다. 소장: 미국 의회도서관, 자료 번호 2003680406.',14.8)
h('한 번의 항해 뒤에 이어진 정복과 교류')
p('이 항해 이후 유럽과 아메리카 사이의 왕래가 지속적으로 확대되었습니다. 유럽인은 새로운 교역과 자원을 추구했고, 스페인 등은 아메리카에서 정복과 식민지 건설을 진행했습니다. 작물과 동물도 대륙 사이를 이동했습니다.')
p('원주민에게 이 변화는 큰 피해를 가져왔습니다. 외부에서 들어온 질병과 전쟁, 강제노동, 토지 상실이 사회를 흔들었습니다. 이 역사는 탐험가의 모험뿐 아니라, 이미 살고 있던 사람들의 삶이 바뀐 역사이기도 합니다. [1, 3, 4]')
h('왜 곧바로 영국 식민지가 된 것은 아닐까요')
p('아메리카 전체가 한 나라의 식민지가 된 적은 없습니다. 스페인, 포르투갈, 프랑스, 네덜란드, 잉글랜드 등이 서로 다른 지역으로 진출했습니다. 미국의 초기 역사를 이해하려면 그중 북아메리카 대서양 연안의 영국 식민지에 초점을 맞춰야 합니다. [4]')

page('3  영국 식민지는 어떻게 시작됐을까요')
p('잉글랜드는 무역과 자원 확보, 다른 유럽 국가와의 경쟁 때문에 북아메리카 진출을 추진했습니다. 왕실이 회사나 개인에게 식민지 건설을 허가하는 특허장을 주고, 투자자가 돈을 모아 이주민과 물자를 보내는 방식이 활용됐습니다. [4, 5]')
h('로어노크의 실패를 지나 제임스타운으로')
p('1580년대 로어노크에서 정착을 시도했지만 안정적인 식민지로 이어지지 못했습니다. 특히 1587년 정착민들은 1590년 보급대가 돌아왔을 때 사라진 상태였습니다. 이후의 행방은 확실하게 밝혀지지 않았습니다. [6]')
p('1607년 버지니아 회사가 보낸 이주민들은 제임스타운을 세웠습니다. 이곳은 오늘날 미국으로 이어지는 지역에서 최초로 지속된 잉글랜드 정착지가 되었습니다. 처음부터 독립국 미국을 만들려던 것이 아니라, 국왕의 권위 아래 수익을 얻으려는 사업이었습니다. [5, 7]')
pic('jamestown.jpg','그림 3  제임스타운에 도착한 잉글랜드인들을 묘사한 재현 삽화. 당시 사진이나 현장 기록화가 아닙니다. 출처: 미국 국립공원관리청 NPS.',15.7)
h('살아남는 문제와 돈을 버는 문제')
p('정착민들은 질병과 식량 부족에 시달렸습니다. 포우하탄 원주민의 지원과 교역은 초기 생존에 중요했지만, 토지와 식량을 둘러싼 갈등은 전쟁으로도 이어졌습니다. 이후 담배 수출이 수익원이 되면서 농장과 정착지가 확대됐습니다. [5, 7]')
p('1624년에는 버지니아 회사의 특허장이 취소되고 버지니아가 왕실 식민지가 됐습니다. 식민지는 세워질 때부터 모두 같은 방식으로 운영된 것이 아니라, 회사의 운영에서 왕실의 직접 통치로 바뀌기도 했습니다. [7]')

page('4  정착지가 13개 식민지로 늘어나다')
h('플리머스에는 종교적 이유로 온 사람들도 있었습니다')
p('1620년 메이플라워호의 이주민들은 플리머스에 정착했습니다. 일부는 자신들의 신앙 공동체를 지키려는 분리파 신자였고, 다른 이유로 참여한 사람들도 있었습니다. 모두가 같은 종교적 목적을 가진 것은 아닙니다. [8]')
p('이들은 공동체 운영을 위한 메이플라워 서약을 맺었습니다. 이듬해에는 왐파노아그 지도자와 동맹을 맺고 원주민의 농사 지식 등 도움을 받았습니다. 그러나 초기 협력만으로 이후의 역사를 설명할 수는 없습니다. 정착지 확대는 원주민의 토지 상실과 충돌로도 이어졌습니다. [4, 8]')
h('동부 해안에서 서로 다른 사회가 성장했습니다')
p('17세기와 18세기에 식민지가 계속 생기고 재편되면서, 독립 당시에는 13개 식민지가 함께 움직였습니다. 이 지역은 오늘날 미국 전체가 아니라 대서양 연안에 집중되어 있었습니다. [4, 9]')
table(['지역','13개 식민지','주요 경제 활동'],[
('뉴잉글랜드','매사추세츠, 뉴햄프셔,\n로드아일랜드, 코네티컷','농업, 어업, 조선, 무역'),
('중부','뉴욕, 뉴저지,\n펜실베이니아, 델라웨어','곡물 농업, 항구 무역'),
('남부','메릴랜드, 버지니아,\n노스캐롤라이나, 사우스캐롤라이나, 조지아','담배·쌀 등 상품 작물 재배')],[2.6,8.5,5.6])
p('표 1  독립 당시 13개 식민지의 일반적인 지역 구분. 플리머스는 1691년 매사추세츠에 통합되어 별도의 14번째 식민지로 세지 않습니다. [9]','Caption')
h('성장의 이면에는 강제노동이 있었습니다')
p('농장과 무역의 성장은 많은 노동력을 요구했습니다. 계약 기간 동안 일하는 이주 노동자도 있었지만, 아프리카인을 강제로 데려와 노동시키는 노예제도 확대됐습니다. 식민지의 경제 성장과 자유의 역사를 말할 때 이들의 강제노동과 원주민의 희생도 함께 살펴야 합니다. [4, 7]')
p('이름 알아두기  초기 정착 당시 국가는 정확히는 잉글랜드입니다. 1707년 잉글랜드와 스코틀랜드가 통합되어 그레이트브리튼이 되었으므로, 앞뒤 시기를 묶어 설명할 때 흔히 “영국 식민지”라고 부릅니다. [10]')

page('5  식민지에서 미국이라는 나라로')
p('식민지 주민들은 오랫동안 영국 국왕의 신민으로 살았습니다. 그러나 1763년 전쟁이 끝난 뒤 영국이 식민지에 세금을 부과하고 통제를 강화하자 갈등이 커졌습니다. 식민지 측은 자신들의 대표가 없는 영국 의회가 세금을 정하는 것에 반발했습니다. [11]')
p('1775년 무력 충돌이 시작됐고, 1776년 7월 4일 대륙회의는 독립선언서를 채택했습니다. 이것이 미국의 탄생을 설명할 때 기준으로 삼는 사건입니다. 다만 선언만으로 전쟁이 끝난 것은 아닙니다. 1783년 파리조약에서 영국이 미국의 독립을 인정했습니다. [11, 12]')
h('시간 순서로 다시 보기')
table(['시기','사건','의미'],[
('유럽인 도착 이전','다양한 원주민 사회','아메리카에는 이미 사람들이 살고 있었습니다.'),
('약 1000년','노르드인의 북아메리카 도착','캐나다에 유적이 남아 있습니다.'),
('1492년','콜럼버스의 카리브해 도착','유럽과 아메리카의 지속적인 접촉이 확대됩니다.'),
('1607년','제임스타운 건설','잉글랜드의 지속적인 식민지 정착이 시작됩니다.'),
('1620년','플리머스 정착','종교와 생계 등 여러 이유로 이주가 이어집니다.'),
('1776년','13개 식민지의 독립선언','미국이라는 새로운 국가를 선언합니다.'),
('1783년','파리조약','영국이 미국의 독립을 인정합니다.')],[3,5,8.7])
p('표 2  아메리카의 오랜 역사와 미국의 국가 형성을 나누어 읽는 연표. [1, 2, 3, 7, 8, 11, 12]','Caption')
h('처음 질문에 대한 답')
p('“미국은 어떻게 발견되고 영국 식민지가 되었을까?”라는 질문은 두 단계로 나누면 이해하기 쉽습니다. 먼저 원주민이 살던 아메리카에 유럽인들이 도착했습니다. 그 뒤 북아메리카 동부에 잉글랜드의 식민지가 만들어지고 성장했으며, 그중 13개가 독립해 미국이 되었습니다.')
p('기억할 핵심  콜럼버스의 도착은 1492년, 제임스타운 건설은 1607년, 미국의 독립선언은 1776년입니다. 서로 다른 사건이며, 그 사이에는 약 300년에 걸친 변화가 있었습니다.')

page('참고 자료와 이미지 출처')
p('본문의 대괄호 번호는 아래 자료와 연결됩니다. 설명은 자료를 바탕으로 한국어로 풀어썼으며, 그림 2와 3은 후대의 재현 이미지입니다. 자료 확인일은 2026년 9월 10일입니다.')
sources=[
('1','미국 의회도서관','What Came to Be Called America','https://www.loc.gov/exhibits/1492/america.html'),
('2','Parks Canada','L’Anse aux Meadows National Historic Site','https://www.pc.gc.ca/apps/dfhd/page_nhs_eng.aspx?id=236'),
('3','미국 의회도서관','Columbus and the Taíno','https://www.loc.gov/exhibits/exploring-the-early-americas/columbus-and-the-taino.html'),
('4','미국 의회도서관','Colonial Settlement 1600s to 1763 Overview','https://www.loc.gov/classroom-materials/united-states-history-primary-source-timeline/colonial-settlement-1600-1763/overview/'),
('5','미국 국립공원관리청','The Virginia Company of London','https://www.nps.gov/jame/learn/historyculture/the-virginia-company-of-london.htm'),
('6','미국 국립공원관리청','1587 The Lost Colony','https://www.nps.gov/fora/learn/historyculture/1587-the-lost-colony.htm'),
('7','미국 국립공원관리청','A Short History of Jamestown · 그림 3 수록','https://www.nps.gov/jame/learn/historyculture/a-short-history-of-jamestown.htm'),
('8','Plimoth Patuxet Museums','Who Were the Pilgrims','https://plimoth.org/for-students/homework-help/who-were-the-pilgrims'),
('9','HISTORY','13 Colonies Map Original States and Facts','https://www.history.com/articles/thirteen-colonies'),
('10','영국 의회','The Articles constitution and trade','https://www.parliament.uk/about/living-heritage/evolutionofparliament/legislativescrutiny/act-of-union-1707/overview/the-articles-constitution-and-trade/'),
('11','미국 국립문서기록관리청','The Declaration of Independence A History','https://www.archives.gov/founding-docs/declaration-history'),
('12','미국 국립문서기록관리청','Treaty of Paris 1783','https://www.archives.gov/milestone-documents/treaty-of-paris'),
('그림 2','미국 의회도서관','Columbus landing on Hispaniola · 자료 2003680406','https://www.loc.gov/pictures/resource/cph.3b07443/'),
('지도 바탕','Natural Earth','1 대 1억 1천만 육지 데이터 · public domain','https://www.naturalearthdata.com/about/terms-of-use/')]
for num,org,title,url in sources:
    pp=p(f'[{num}] {org} | {title}');pp.paragraph_format.space_after=Pt(2)
    for r in pp.runs:r.font.size=Pt(8.5);r.bold=True
    pp=p(url);pp.paragraph_format.space_after=Pt(7);pp.paragraph_format.line_spacing=1.0
    for r in pp.runs:r.font.size=Pt(7.5)
doc.core_properties.title='아메리카와 미국의 시작';doc.core_properties.subject='유럽인의 도착과 영국 식민지 형성을 설명하는 한국어 학습 자료';doc.core_properties.author=''
doc.save(ROOT/'아메리카와_미국의_시작.docx')
print('Created',ROOT/'아메리카와_미국의_시작.docx')
